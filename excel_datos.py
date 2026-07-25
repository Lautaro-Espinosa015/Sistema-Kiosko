"""
Módulo que se encarga de TODO lo relacionado al archivo Excel:
leerlo al iniciar la app y escribirlo al guardar.

Ni Inventario ni Caja saben que el Excel existe: este módulo
las "traduce" a filas de una planilla y viceversa.
"""

import os
import sys
import shutil
from datetime import datetime
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

from caja import Caja
from inventario import Inventario
from modelos import Venta

NOMBRE_ARCHIVO = "kiosko_datos.xlsx"


def _directorio_base() -> str:
    """
    Devuelve la carpeta donde vive el programa.
    - Si es un .exe armado con PyInstaller, es la carpeta del .exe.
    - Si es el script .py normal, es la carpeta donde está este archivo.
    """
    if getattr(sys, "frozen", False):
        return os.path.dirname(sys.executable)
    return os.path.dirname(os.path.abspath(__file__))


RUTA_ARCHIVO = os.path.join(_directorio_base(), NOMBRE_ARCHIVO)

COLUMNAS_PRODUCTOS = ["Código", "Nombre", "Categoría", "Tipo", "Precio", "Costo", "Stock"]
COLUMNAS_VENTAS = [
    "Código", "Nombre", "Categoría", "Tipo", "Cantidad", "Precio unitario",
    "Costo unitario", "Total", "Ganancia", "Forma de pago",
]
COLUMNAS_RESUMEN = ["Concepto", "Valor"]


def cargar_datos(inventario: Inventario, caja: Caja) -> bool:
    """
    Carga productos y ventas del Excel hacia el inventario y la caja.
    Devuelve True si el archivo ya existía, o False si no existía
    y se creó uno nuevo (vacío) en su lugar.
    """
    if not os.path.exists(RUTA_ARCHIVO):
        _crear_archivo_vacio()
        return False

    libro = load_workbook(RUTA_ARCHIVO, data_only=True)

    if "Productos" in libro.sheetnames:
        hoja = libro["Productos"]
        for fila in hoja.iter_rows(min_row=2, values_only=True):
            if fila[0] is None:
                continue
            codigo, nombre, categoria, tipo, precio, costo, stock = fila[:7]
            inventario.agregar_producto(
                int(codigo),
                str(nombre),
                str(categoria or ""),
                str(tipo or "Unidad"),
                float(precio),
                float(costo),
                float(stock),
            )

    if "Ventas" in libro.sheetnames:
        hoja = libro["Ventas"]
        for fila in hoja.iter_rows(min_row=2, values_only=True):
            if fila[0] is None:
                continue
            codigo, nombre, categoria, tipo, cantidad, precio_unitario, costo_unitario = fila[:7]
            forma_pago = fila[9] if len(fila) > 9 and fila[9] else "Efectivo"
            venta = Venta(
                codigo=int(codigo),
                nombre=str(nombre),
                categoria=str(categoria or ""),
                tipo=str(tipo or "Unidad"),
                cantidad=float(cantidad),
                precio_unitario=float(precio_unitario),
                costo_unitario=float(costo_unitario),
                forma_pago=str(forma_pago),
            )
            caja.agregar_venta_existente(venta)

    return True


def guardar_datos(inventario: Inventario, caja: Caja) -> None:
    """Guarda el estado actual (productos y ventas) en el Excel, pisando lo anterior."""
    libro = Workbook()

    hoja_productos = libro.active
    hoja_productos.title = "Productos"
    _escribir_encabezado(hoja_productos, COLUMNAS_PRODUCTOS)
    for producto in inventario.obtener_todos():
        hoja_productos.append([
            producto.codigo, producto.nombre, producto.categoria, producto.tipo,
            producto.precio, producto.costo, producto.stock,
        ])

    hoja_ventas = libro.create_sheet("Ventas")
    _escribir_encabezado(hoja_ventas, COLUMNAS_VENTAS)
    for venta in caja.obtener_ventas():
        hoja_ventas.append([
            venta.codigo, venta.nombre, venta.categoria, venta.tipo, venta.cantidad,
            venta.precio_unitario, venta.costo_unitario,
            venta.total, venta.ganancia, venta.forma_pago,
        ])

    hoja_resumen = libro.create_sheet("Resumen")
    _escribir_encabezado(hoja_resumen, COLUMNAS_RESUMEN)
    hoja_resumen.append(["Cantidad de ventas", caja.cantidad_operaciones()])
    hoja_resumen.append(["Total vendido", caja.total_vendido()])
    hoja_resumen.append(["Ganancia total", caja.total_ganancia()])
    for forma_pago, total in caja.total_por_forma_pago().items():
        hoja_resumen.append([f"Total en {forma_pago}", total])

    for hoja in (hoja_productos, hoja_ventas, hoja_resumen):
        _autoajustar_columnas(hoja)

    libro.save(RUTA_ARCHIVO)


def _crear_archivo_vacio() -> None:
    """Crea kiosko_datos.xlsx con las hojas y encabezados, sin filas de datos."""
    libro = Workbook()

    hoja_productos = libro.active
    hoja_productos.title = "Productos"
    _escribir_encabezado(hoja_productos, COLUMNAS_PRODUCTOS)

    hoja_ventas = libro.create_sheet("Ventas")
    _escribir_encabezado(hoja_ventas, COLUMNAS_VENTAS)

    hoja_resumen = libro.create_sheet("Resumen")
    _escribir_encabezado(hoja_resumen, COLUMNAS_RESUMEN)

    libro.save(RUTA_ARCHIVO)


def _escribir_encabezado(hoja, columnas: list[str]) -> None:
    for indice, titulo in enumerate(columnas, start=1):
        celda = hoja.cell(row=1, column=indice, value=titulo)
        celda.font = Font(bold=True)


def _autoajustar_columnas(hoja) -> None:
    """Ensancha cada columna según el contenido más largo, para que se lea bien."""
    for columna in hoja.columns:
        letra = columna[0].column_letter
        largo_max = max((len(str(c.value)) for c in columna if c.value is not None), default=8)
        hoja.column_dimensions[letra].width = largo_max + 2
        
def archivar_periodo_y_reiniciar(inventario: Inventario, caja: Caja, etiqueta_periodo: str) -> str:
    """
    1. Hace una copia de respaldo del Excel actual etiquetado con la fecha/nombre del periodo.
    2. Reinicia la lista de ventas en caja.
    3. Sobrescribe 'kiosko_datos.xlsx' dejando Ventas en cero pero manteniendo los Productos/Stock.
     Devuelve el nombre del archivo de respaldo creado.
    """
    # Guardamos los datos actuales para asegurarnos de que el archivo base esté al día
    guardar_datos(inventario, caja)
    
    # Nombre para el archivo de respaldo/histórico
    timestamp = datetime.now().strftime("%Y-%m-%d_%H-%M")
    nombre_archivo_backup = f"kiosko_ventas_{etiqueta_periodo}_{timestamp}.xlsx"
    ruta_backup = os.path.join(_directorio_base(), nombre_archivo_backup)
    
    # Copiamos el excel actual con las ventas cerradas
    shutil.copyfile(RUTA_ARCHIVO, ruta_backup)
    
    # Vaciamos la caja en memoria
    caja.limpiar_ventas()
    
    # Sobrescribimos el excel activo con la caja limpia (Ventas vacías, Productos intactos)
    guardar_datos(inventario, caja)
    
    return nombre_archivo_backup